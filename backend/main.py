"""FastAPI 入口：健康检查、错题 CRUD、AI 识图 / 修正、配置读写。

本地运行在 http://127.0.0.1:8000，供 Electron 主进程 spawn 拉起。
所有用户数据（DB / uploads / config）统一落在 USER_DATA 目录（见 database.py）。
"""

import base64
import hashlib
import json
import os
import re
import sys
from contextlib import asynccontextmanager
from datetime import date
from pathlib import Path
from typing import List, Optional

# PyInstaller 打包运行时：conda 扩展模块（_sqlite3 等）依赖的 DLL 在
# 解压目录 _MEIPASS 里，需显式加入 DLL 搜索路径（见 backend.spec）。
if getattr(sys, "frozen", False) and sys.platform == "win32":
    os.add_dll_directory(sys._MEIPASS)

import aiofiles
import uvicorn
from fastapi import FastAPI, File, HTTPException, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware

import config
import database
from models import (
    ConfigModel,
    CorrectLatexRequest,
    ExportRequest,
    GenerateAnswerRequest,
    GenerateSimilarRequest,
    ProblemCreate,
    ProblemOut,
    ProblemUpdate,
)
from services import ai_client, effort, exporter, scheduler, textfix

# 支持的图片扩展名 → MIME（走多模态 vision）。
IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
}

# 语言一致性规则：AI 输出必须与题目原文同语言（英文题不要翻译成中文，反之亦然）。
# 所有 AI 调用（识图 / 文本整理 / 生成答案 / 举一反三 / 修正）共用。
LANGUAGE_RULE = (
    "5. 【语言一致性】输出内容的语言必须与题目原文的语言完全一致："
    "题目是英文就全程用英文（题干、解答、标签都用英文），"
    "题目是中文就全程用中文；题目为其它语言时同样沿用该语言。"
    "严禁把题目翻译成另一种语言，也严禁中英混杂。\n"
)

# ---------------------------------------------------------------------------
# 提示词：所有 AI 调用点共用一份格式契约
# ---------------------------------------------------------------------------
# 【为什么必须统一】前端只有一套渲染方式（KaTeX 扫 \( \) 与 \[ \]），
# 所以所有 AI 产出的文本都必须是同一种格式。曾经分成两套——识图/举一反三要 JSON、
# 生成答案要纯文本——结果模型按 JSON 那套输出答案，库里存进了整坨
# {"latex_code":"..."}，前端渲染出一堆源码。
#
# 现在的做法：**所有调用点都要 JSON 信封**，都用 textfix.parse_ai_json 解析。
# 理由是模型天然倾向返回 JSON（拦不住），而 JSON 的转义规则是确定的，
# 配上 textfix 的非法转义修复就能稳定拿到正文；反过来「求模型别用 JSON」是碰运气。
FORMAT_CONTRACT = (
    "严格遵守以下格式规则：\n"
    "1. 所有数学/物理/化学公式必须用 \\\\( \\\\) 包裹行内公式、用 \\\\[ \\\\] 包裹行间公式；"
    "禁止使用 $ … $、$$ … $$、\\\\begin{equation}、\\\\begin{align} 等任何其它公式定界符。\n"
    "2. 内容里有多个部分（多道题、多个小问、多个解题步骤）时，每部分各占一段，"
    "段与段之间用换行符 \\n 分隔，并保留原有编号（如 1.、2.、(a)、(i)）。\n"
    "3. 返回值是 JSON 字符串，LaTeX 里的反斜杠必须按 JSON 规则转义："
    "\\\\frac 要写成 \\\\\\\\frac，\\\\( 要写成 \\\\\\\\(。"
    "这一条最容易出错——反斜杠没转义会导致 JSON 解析失败、公式无法显示。\n"
    "4. 只返回一个 JSON 对象，不要输出任何解释、前后缀或 Markdown 代码块标记。\n"
    + LANGUAGE_RULE
)

# 录入类调用（识图 / 文本整理）的字段说明。
JSON_SPEC = (
    "JSON 字段如下：\n"
    '{"subject": "数学/物理/化学 之一", '
    '"latex_code": "题目内容（遵守上述格式规则）", '
    '"raw_text": "题目的纯文本形式", '
    '"tags": ["知识点标签1", "知识点标签2", "知识点标签3"], '
    '"answer_latex": "题目的解答/答案（遵守上述格式规则；若原图或原文中没有解答，则输出空字符串 \\"\\"）"}'
    "\n"
    "tags 字段是必填项，不得为空数组：必须为每道题给出 2-4 个具体的知识点标签"
    "（例如「二次函数」「数列递推」「牛顿第二定律」「氧化还原反应」这种细分知识点，"
    "而不是「数学」「物理」这类学科名）。标签语言与题目语言保持一致。\n"
    "注意：subject 是固定枚举，只能填「数学」「物理」「化学」三个中文词之一，"
    "不受上述语言一致性规则约束——即使题目是英文，subject 也必须是中文。\n"
)

# 单字段调用（生成答案 / 修正题目）的字段说明。
# 同样是 JSON 信封，与录入类保持一致，只是字段少。
ANSWER_SPEC = (
    "JSON 字段如下：\n"
    '{"answer_latex": "完整解答（遵守上述格式规则）"}\n'
)

QUESTION_SPEC = (
    "JSON 字段如下：\n"
    '{"latex_code": "修正后的完整题目内容（遵守上述格式规则）"}\n'
)

OCR_PROMPT = (
    "你是一个数理化错题识别助手。请识别图片中的全部题目。\n" + FORMAT_CONTRACT + JSON_SPEC
)

TEXT_PROMPT = (
    "你是一个数理化错题整理助手。请整理下面从文件中提取的题目文本。\n"
    + FORMAT_CONTRACT
    + JSON_SPEC
    + "\n\n提取到的文本如下：\n"
)


@asynccontextmanager
async def lifespan(app: FastAPI):
    """应用生命周期：启动时初始化数据库。"""
    database.init_db()
    yield


app = FastAPI(title="数理化错题本 API", version="0.2.0", lifespan=lifespan)

# 允许前端（Electron file:// 或 localhost）跨域访问本地 API。
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------
# 统一的查询列：附带 seq —— 按 id 升序的连续序号。
# 用子查询（COUNT id <= 自身）而不是 rowid，保证删除错题后序号依然连续，
# 前端展示 seq、接口调用仍用 id。
SELECT_COLS = (
    "p.*, (SELECT COUNT(*) FROM problems q WHERE q.id <= p.id) AS seq"
)


# 学科在库里统一存中文枚举（数学/物理/化学），界面再按语言翻译显示。
# 但「输出跟随题目语言」的规则会让模型对英文题返回 "Mathematics"，
# 导致学科色块失效、自动附加的学科标签中英分裂（同一学科两种标签、
# 按标签搜索被割裂）。故入库前统一归一化。
SUBJECT_ALIASES = {
    "math": "数学",
    "maths": "数学",
    "mathematics": "数学",
    "physics": "物理",
    "physic": "物理",
    "chemistry": "化学",
    "chem": "化学",
}


def _normalize_subject(subject: Optional[str]) -> Optional[str]:
    """把英文/异体学科名归一成中文枚举；无法识别时原样返回。"""
    if not subject:
        return subject
    s = subject.strip()
    return SUBJECT_ALIASES.get(s.lower(), s)


# 学科枚举（库里的合法取值）；同步学科标签时用它判断哪些标签是「学科标签」。
SUBJECT_ENUM = ("数学", "物理", "化学")


def _sync_subject_tag(
    tags_json: Optional[str], old_subject: Optional[str], new_subject: Optional[str]
) -> List[str]:
    """改学科时同步学科标签：摘掉旧学科标签，把新学科放到最前面。

    只动学科枚举内的标签（数学/物理/化学）和旧学科本身，用户自定义的
    知识点标签一律保留、顺序不变。
    """
    try:
        tags = json.loads(tags_json) if tags_json else []
    except (json.JSONDecodeError, TypeError):
        tags = []
    if not isinstance(tags, list):
        tags = []

    drop = set(SUBJECT_ENUM)
    if old_subject:
        drop.add(old_subject)
    kept = [tag for tag in tags if isinstance(tag, str) and tag not in drop]
    return ([new_subject] + kept) if new_subject else kept


def _split_q_a(content: str) -> tuple[str, Optional[str]]:
    """把「【题目】…【解答】…」的内容拆成（题目, 答案）。

    旧数据里题目和解答混在 latex_code 中（如举一反三生成的题），
    这里按标记拆分；没有标记则整段视为题目，答案为 None。
    """
    if not content:
        return "", None
    m = re.search(r"【\s*题目\s*】(.*?)【\s*解答\s*】(.*)", content, re.DOTALL)
    if m:
        question = m.group(1).strip()
        answer = m.group(2).strip()
        return question, answer or None
    return content.strip(), None


def _row_to_problem(row) -> ProblemOut:
    """把 sqlite3.Row 转成 ProblemOut，tags 字段从 JSON 字符串反序列化。

    question_latex / answer_latex：优先用库里的 answer_latex 列；
    为空时尝试从 latex_code 按【题目】/【解答】标记拆分（兼容旧数据）。
    """
    raw_tags = row["tags"]
    try:
        tags = json.loads(raw_tags) if raw_tags else []
    except (json.JSONDecodeError, TypeError):
        tags = []
    question, split_answer = _split_q_a(row["latex_code"] or row["raw_text"] or "")
    answer = row["answer_latex"] or split_answer
    keys = row.keys()
    return ProblemOut(
        id=row["id"],
        seq=row["seq"] if "seq" in keys else 0,
        image_path=row["image_path"],
        raw_text=row["raw_text"],
        latex_code=row["latex_code"],
        question_latex=question or None,
        answer_latex=answer or None,
        subject=row["subject"],
        tags=tags,
        created_at=row["created_at"],
        next_review_date=row["next_review_date"],
        review_stage=row["review_stage"],
        raw_image_hash=row["raw_image_hash"],
        is_generated=bool(row["is_generated"]),
        parent_id=row["parent_id"],
        last_review_date=row["last_review_date"],
    )


def _fetch_problem(conn, problem_id: int):
    """按 id 查一条错题（含 seq 连续序号），返回 sqlite3.Row 或 None。"""
    return conn.execute(
        f"SELECT {SELECT_COLS} FROM problems p WHERE p.id = ?", (problem_id,)
    ).fetchone()


def _parse_ai_json(text: str) -> dict:
    """解析录入类调用返回的 JSON（识图 / 文本整理 / 举一反三共用）。

    统一走 textfix.parse_ai_json：它会先修 LaTeX 造成的非法 JSON 转义
    （`"\\("` 不是合法转义，模型经常这么写，不修就整段解析失败），
    再按需做截断兜底。彻底拿不到才退回「整段文本当题目」。
    """
    data = textfix.parse_ai_json(
        text, keys=("latex_code", "answer_latex", "raw_text")
    )
    if data:
        return data
    return {"raw_text": text, "latex_code": text, "subject": None, "tags": []}


# 文本收敛与定界符统一都放在 services/textfix.py（database.py 的老数据迁移也要用，
# 放这里会循环导入）。这两个别名保持原有调用点不动。
_coerce_ai_text = textfix.coerce_ai_text
_sanitize_latex = textfix.sanitize_latex


def _extract_pdf_text(path: Path) -> str:
    """提取 PDF 文本层（扫描版纯图片 PDF 会得到空串）。"""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx_text(path: Path) -> str:
    """提取 DOCX 段落文本。"""
    import docx

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs).strip()


def _insert_problem(
    *,
    image_path: Optional[str],
    raw_text: Optional[str],
    latex_code: Optional[str],
    subject: Optional[str],
    tags: List[str],
    raw_image_hash: Optional[str],
    is_generated: bool = False,
    parent_id: Optional[int] = None,
    answer_latex: Optional[str] = None,
) -> ProblemOut:
    """插入一条错题：created_at 为当天，next_review_date 按艾宾浩斯首个间隔计算。

    tags 自动带上学科标签（如数学/物理/化学），方便按学科筛选。
    """
    today = date.today().isoformat()
    # 学科归一化：模型对英文题可能返回 "Mathematics"，统一转中文枚举，
    # 保证色块、自动标签、按标签搜索在中英界面下都一致。
    subject = _normalize_subject(subject)
    # 学科标签自动附加（去重）：所有录入路径统一生效。
    tags = list(tags or [])
    if subject and subject not in tags:
        tags = [subject] + tags
    tags_json = json.dumps(tags, ensure_ascii=False)
    next_review = scheduler.next_review_date(0)

    conn = database.get_connection()
    try:
        cursor = conn.execute(
            """
            INSERT INTO problems
                (image_path, raw_text, latex_code, subject, tags,
                 created_at, next_review_date, review_stage, raw_image_hash,
                 is_generated, parent_id, answer_latex)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                image_path,
                raw_text,
                latex_code,
                subject,
                tags_json,
                today,
                next_review,
                0,
                raw_image_hash,
                1 if is_generated else 0,
                parent_id,
                answer_latex,
            ),
        )
        conn.commit()
        row = _fetch_problem(conn, cursor.lastrowid)
    finally:
        conn.close()

    if row is None:
        raise HTTPException(status_code=500, detail="创建错题后未能读取记录")
    return _row_to_problem(row)


# ---------------------------------------------------------------------------
# 基础接口
# ---------------------------------------------------------------------------
@app.get("/health")
def health() -> dict:
    """健康检查：返回服务运行状态。"""
    return {"status": "ok"}


@app.get("/api/problems", response_model=List[ProblemOut])
def list_problems() -> List[ProblemOut]:
    """获取错题列表（按创建时间倒序）。"""
    conn = database.get_connection()
    try:
        rows = conn.execute(
            f"SELECT {SELECT_COLS} FROM problems p ORDER BY p.id DESC"
        ).fetchall()
    finally:
        conn.close()
    return [_row_to_problem(row) for row in rows]


@app.post("/api/problems", response_model=ProblemOut, status_code=201)
def create_problem(problem: ProblemCreate) -> ProblemOut:
    """新增错题条目（手动录入）。"""
    return _insert_problem(
        image_path=problem.image_path,
        raw_text=problem.raw_text,
        latex_code=problem.latex_code,
        subject=problem.subject,
        tags=problem.tags,
        raw_image_hash=None,
    )


@app.get("/api/problems/{problem_id}", response_model=ProblemOut)
def get_problem(problem_id: int) -> ProblemOut:
    """按 id 获取单道错题。"""
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, problem_id)
    finally:
        conn.close()
    if row is None:
        raise HTTPException(status_code=404, detail="错题不存在")
    return _row_to_problem(row)


@app.patch("/api/problems/{problem_id}", response_model=ProblemOut)
def update_problem(problem_id: int, patch: ProblemUpdate) -> ProblemOut:
    """手动编辑错题（部分字段），支撑前端 LaTeX 源码「更新」按钮。"""
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, problem_id)
        if row is None:
            raise HTTPException(status_code=404, detail="错题不存在")

        fields = []
        values = []
        if patch.raw_text is not None:
            fields.append("raw_text = ?")
            values.append(patch.raw_text)
        if patch.latex_code is not None:
            fields.append("latex_code = ?")
            values.append(patch.latex_code)
        if patch.answer_latex is not None:
            fields.append("answer_latex = ?")
            values.append(patch.answer_latex)

        # 学科：归一化成中文枚举（英文界面改学科时前端也可能传英文），
        # 空串视为「清空学科」→ NULL。
        new_subject = None
        if patch.subject is not None:
            new_subject = _normalize_subject(patch.subject) or None
            fields.append("subject = ?")
            values.append(new_subject)

        # 标签：调用方显式传 tags 时以它为准；否则改学科要顺带同步学科标签
        # （录入时自动附加过学科标签，不同步会留下旧学科的标签，
        #  导致按标签搜索把同一道题算进旧学科）。
        if patch.tags is not None:
            fields.append("tags = ?")
            values.append(json.dumps(patch.tags, ensure_ascii=False))
        elif patch.subject is not None:
            synced = _sync_subject_tag(row["tags"], row["subject"], new_subject)
            fields.append("tags = ?")
            values.append(json.dumps(synced, ensure_ascii=False))

        if fields:
            values.append(problem_id)
            conn.execute(
                f"UPDATE problems SET {', '.join(fields)} WHERE id = ?", values
            )
            conn.commit()

        row = _fetch_problem(conn, problem_id)
    finally:
        conn.close()
    return _row_to_problem(row)


@app.delete("/api/problems/{problem_id}", status_code=204)
def delete_problem(problem_id: int) -> Response:
    """删除错题，并清理其孤立的上传文件（每个文件按 hash 唯一，删记录即可删文件）。"""
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, problem_id)
        if row is None:
            raise HTTPException(status_code=404, detail="错题不存在")
        image_path = row["image_path"]
        conn.execute("DELETE FROM problems WHERE id = ?", (problem_id,))
        # 全部删空时重置自增计数器，下次录入从 id=1 重新开始（避免 id 无限增长）。
        remaining = conn.execute("SELECT COUNT(*) FROM problems").fetchone()[0]
        if remaining == 0:
            conn.execute("DELETE FROM sqlite_sequence WHERE name = 'problems'")
        conn.commit()
    finally:
        conn.close()

    if image_path:
        try:
            p = Path(image_path)
            if p.exists():
                p.unlink()
        except OSError:
            pass  # 文件清理失败不影响删除结果

    return Response(status_code=204)


# ---------------------------------------------------------------------------
# 配置接口
# ---------------------------------------------------------------------------
@app.get("/api/config", response_model=ConfigModel)
def get_config() -> ConfigModel:
    """读取 AI 配置。"""
    return ConfigModel(**config.read_config())


@app.post("/api/config", response_model=ConfigModel)
def save_config(cfg: ConfigModel) -> ConfigModel:
    """保存 AI 配置。"""
    saved = config.write_config(cfg.model_dump())
    return ConfigModel(**saved)


@app.get("/api/effort/levels")
def effort_levels(
    model: Optional[str] = None, requested: Optional[str] = None
) -> dict:
    """返回推理强度档位信息，供「设置」页构建下拉框。

    model     ：缺省时用当前配置里的模型。
    requested ：缺省时用配置里已存的档位。传入下拉框里**当前选中**的档位，
                就能在用户还没点保存时就告诉他「这一档会被降级成 X」。

    返回的 source 说明可信度：
    probe（真实探测过）> learned（调用中撞过墙）> catalog（已知表）> default（未知）。
    effective 是实际会发给上游的档位（auto 表示不发送该参数），
    由 effort.resolve() 统一计算——前端不再自己复制一份降级规则。
    """
    cfg = config.read_config()
    target = (model or cfg.get("model_name") or "").strip()
    info = effort.describe(cfg.get("base_url", ""), target)
    want = requested if requested is not None else cfg.get("reasoning_effort", "auto")
    supported = None if info["source"] == "default" else info["supported"]
    info["requested"] = want
    info["effective"] = effort.resolve(want, supported) or effort.AUTO
    return info


@app.post("/api/effort/probe")
async def effort_probe() -> dict:
    """真实探测当前模型支持哪些档位（每档发一次极小请求），结果写入缓存。"""
    try:
        return await ai_client.probe_efforts()
    except ai_client.AIConfigError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ai_client.AIRequestError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.get("/api/models")
async def list_upstream_models() -> dict:
    """拉上游可用模型列表（GET {base_url}/models），供「设置」页下拉选择。"""
    try:
        return {"models": await ai_client.list_models()}
    except ai_client.AIConfigError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ai_client.AIRequestError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


# ---------------------------------------------------------------------------
# AI 接口
# ---------------------------------------------------------------------------
@app.post("/api/upload", response_model=ProblemOut, status_code=201)
async def upload_problem(file: UploadFile = File(...)) -> ProblemOut:
    """上传图片/PDF/DOCX，AI 识别后存库并返回。

    图片走多模态 vision；PDF/DOCX 提取文本层后交给 AI 结构化。
    以文件内容 sha256 去重：命中已存在记录则直接返回旧记录。
    """
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    file_hash = hashlib.sha256(content).hexdigest()

    # 去重：同一文件已上传过则直接返回旧记录。
    conn = database.get_connection()
    try:
        existing = conn.execute(
            f"SELECT {SELECT_COLS} FROM problems p WHERE p.raw_image_hash = ?",
            (file_hash,),
        ).fetchone()
    finally:
        conn.close()
    if existing is not None:
        return _row_to_problem(existing)

    ext = Path(file.filename or "").suffix.lower()
    saved_path = database.get_uploads_dir() / f"{file_hash}{ext}"
    async with aiofiles.open(saved_path, "wb") as f:
        await f.write(content)

    try:
        if ext in IMAGE_MIME:
            image_b64 = base64.b64encode(content).decode("ascii")
            ai_text = await ai_client.call_ai(
                OCR_PROMPT, image_base64=image_b64, image_mime=IMAGE_MIME[ext]
            )
        elif ext == ".pdf":
            extracted = _extract_pdf_text(saved_path)
            if not extracted:
                raise HTTPException(
                    status_code=422,
                    detail="PDF 未提取到文本层（可能是扫描版纯图片 PDF），本阶段暂不支持。",
                )
            ai_text = await ai_client.call_ai(TEXT_PROMPT + extracted)
        elif ext == ".docx":
            extracted = _extract_docx_text(saved_path)
            if not extracted:
                raise HTTPException(status_code=422, detail="DOCX 未提取到文本内容。")
            ai_text = await ai_client.call_ai(TEXT_PROMPT + extracted)
        else:
            raise HTTPException(
                status_code=415,
                detail=f"不支持的文件类型：{ext or '未知'}（支持 PNG/JPG/PDF/DOCX）",
            )
    except ai_client.AIConfigError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ai_client.AIRequestError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    parsed = _parse_ai_json(ai_text)
    return _insert_problem(
        image_path=str(saved_path),
        raw_text=parsed.get("raw_text"),
        latex_code=parsed.get("latex_code"),
        subject=parsed.get("subject"),
        tags=parsed.get("tags") or [],
        raw_image_hash=file_hash,
        answer_latex=_sanitize_latex(parsed.get("answer_latex") or "") or None,
    )


@app.post("/api/correct-latex", response_model=ProblemOut)
async def correct_latex(req: CorrectLatexRequest) -> ProblemOut:
    """按用户的自然语言指令修正题目的 LaTeX，并更新数据库。"""
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, req.problem_id)
    finally:
        conn.close()
    if row is None:
        raise HTTPException(status_code=404, detail="错题不存在")

    current_latex = row["latex_code"] or row["raw_text"] or ""
    prompt = (
        "你是一个 LaTeX 题目修正助手。下面是当前题目的内容，请根据用户的修改要求"
        "输出修正后的完整内容。\n"
        + FORMAT_CONTRACT
        + QUESTION_SPEC
        + f"\n当前内容：\n{current_latex}\n\n"
        + f"用户修改要求：{req.user_feedback}"
    )

    try:
        corrected = await ai_client.call_ai(prompt)
    except ai_client.AIConfigError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ai_client.AIRequestError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    # 与录入/生成答案同一个解析器。
    parsed = textfix.parse_ai_json(corrected, keys=("latex_code", "answer_latex"))
    fixed = parsed.get("latex_code") or ""
    if not isinstance(fixed, str) or not fixed.strip():
        fixed = _coerce_ai_text(corrected)
    corrected = _sanitize_latex(fixed)

    conn = database.get_connection()
    try:
        conn.execute(
            "UPDATE problems SET latex_code = ? WHERE id = ?",
            (corrected, req.problem_id),
        )
        conn.commit()
        row = _fetch_problem(conn, req.problem_id)
    finally:
        conn.close()
    return _row_to_problem(row)


@app.post("/api/generate-answer", response_model=ProblemOut)
async def generate_answer(req: GenerateAnswerRequest) -> ProblemOut:
    """用 AI 根据题目生成完整解答，存入 answer_latex 并返回。"""
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, req.problem_id)
    finally:
        conn.close()
    if row is None:
        raise HTTPException(status_code=404, detail="错题不存在")

    question, split_answer = _split_q_a(row["latex_code"] or row["raw_text"] or "")
    if not question:
        raise HTTPException(status_code=400, detail="题目内容为空，无法生成答案")

    current = (row["answer_latex"] or split_answer or "").strip()
    feedback = (req.user_feedback or "").strip()

    # 三种任务：按用户修正意见改 / 换一版重写 / 首次生成。
    # 格式约束不在这里，统一由下面的 FORMAT_CONTRACT + ANSWER_SPEC 提供。
    if feedback and current:
        task = (
            "下面是一道题和它当前的解答。请按用户的修改要求给出修正后的**完整**解答，"
            "保留正确的部分，只改需要改的地方。\n"
        )
        context = f"当前解答：\n{current}\n\n用户修改要求：{feedback}"
    elif feedback:
        task = (
            "请为下面这道题写出完整、清晰的解答步骤，并额外满足用户的具体要求。\n"
        )
        context = f"用户要求：{feedback}"
    elif req.regenerate and current:
        task = (
            "下面是一道题和它当前的解答。用户对当前解答不满意，请**重新独立解一遍**："
            "换一种思路或更清晰的表述，步骤写得更完整；不要照抄当前解答。\n"
        )
        context = f"当前解答（不满意，仅供参考避免重复）：\n{current}"
    else:
        task = (
            "请为下面这道题写出完整、清晰的解答步骤，包含必要的公式推导与最终答案。\n"
        )
        context = ""

    # 与识图/举一反三完全同一套格式契约 + JSON 信封，保证渲染格式一致。
    prompt = (
        "你是一个数理化解题助手。"
        + task
        + FORMAT_CONTRACT
        + ANSWER_SPEC
        + f"\n题目：\n{question}\n"
        + (f"\n{context}\n" if context else "")
    )

    try:
        ai_text = await ai_client.call_ai(prompt)
    except ai_client.AIConfigError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ai_client.AIRequestError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    # 与录入路径同一个解析器：修非法转义 → 取字段 → 统一定界符。
    parsed = textfix.parse_ai_json(ai_text, keys=("answer_latex", "latex_code"))
    answer = parsed.get("answer_latex") or parsed.get("latex_code") or ""
    if not isinstance(answer, str) or not answer.strip():
        # 模型没按信封走（少数情况）→ 退回单段正文收敛，仍能拿到可渲染的内容。
        answer = _coerce_ai_text(ai_text)
    answer = _sanitize_latex(answer) or None

    conn = database.get_connection()
    try:
        conn.execute(
            "UPDATE problems SET answer_latex = ? WHERE id = ?",
            (answer, req.problem_id),
        )
        conn.commit()
        row = _fetch_problem(conn, req.problem_id)
    finally:
        conn.close()
    return _row_to_problem(row)


# ---------------------------------------------------------------------------
# 举一反三
# ---------------------------------------------------------------------------
@app.post("/api/generate-similar", response_model=ProblemOut, status_code=201)
async def generate_similar(req: GenerateSimilarRequest) -> ProblemOut:
    """基于父题用 AI 生成「变式」或「拓展」新题，作为独立错题存库并返回。"""
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, req.problem_id)
    finally:
        conn.close()
    if row is None:
        raise HTTPException(status_code=404, detail="错题不存在")

    gen_type = (req.type or "变式").strip()
    if gen_type == "拓展":
        task = (
            "请基于下面这道题，生成一道【进阶拓展题】：在原题知识点基础上适度提高难度、"
            "综合更多知识点或增加推理步骤，保持同一学科。"
        )
    else:
        gen_type = "变式"
        task = (
            "请基于下面这道题，生成一道【同类变式题】：考查相同知识点，"
            "更换数值/情境/表述，难度与原题相当，保持同一学科。"
        )

    source = row["latex_code"] or row["raw_text"] or ""
    subject = row["subject"]
    # 格式要求与识图/生成答案完全一致（FORMAT_CONTRACT 已含反斜杠转义、
    # 定界符、禁代码块等全部约束，这里不再重复叮嘱一遍）。
    prompt = (
        "你是一个数理化出题助手。"
        + task
        + "\n新题必须附带完整解题步骤。\n"
        + FORMAT_CONTRACT
        + JSON_SPEC
        + "\n特别要求：题干写进 latex_code，解答写进 answer_latex，不要把两者混在一个字段里。\n"
        + f"\n原题（学科：{subject or '未知'}）：\n{source}"
    )

    try:
        ai_text = await ai_client.call_ai(prompt)
    except ai_client.AIConfigError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ai_client.AIRequestError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    parsed = _parse_ai_json(ai_text)
    # 题干与解答现在分字段返回（与识图路径一致）。老提示词是把两者塞进
    # latex_code 再靠【题目】/【解答】标记拆分，这里兜一下：若模型仍混在一起，
    # _split_q_a 能把解答摘出来，不至于丢答案。
    question = _sanitize_latex(parsed.get("latex_code") or "")
    answer = _sanitize_latex(parsed.get("answer_latex") or "")
    if not answer:
        question, split_answer = _split_q_a(question)
        answer = split_answer or ""
    return _insert_problem(
        image_path=None,
        raw_text=parsed.get("raw_text"),
        latex_code=question,
        subject=parsed.get("subject") or subject,
        tags=parsed.get("tags") or [],
        raw_image_hash=None,
        is_generated=True,
        parent_id=req.problem_id,
        answer_latex=answer or None,
    )


# ---------------------------------------------------------------------------
# 复习计划（艾宾浩斯）
# ---------------------------------------------------------------------------
@app.get("/api/review/due", response_model=List[ProblemOut])
def review_due() -> List[ProblemOut]:
    """返回今日（含逾期）待复习、且未掌握的错题。"""
    today = date.today().isoformat()
    max_stage = len(scheduler.INTERVALS)
    conn = database.get_connection()
    try:
        rows = conn.execute(
            f"""
            SELECT {SELECT_COLS} FROM problems p
            WHERE p.next_review_date <= ? AND p.review_stage < ?
            ORDER BY p.next_review_date, p.id
            """,
            (today, max_stage),
        ).fetchall()
    finally:
        conn.close()
    return [_row_to_problem(row) for row in rows]


@app.post("/api/review/{problem_id}/complete", response_model=ProblemOut)
def complete_review(problem_id: int) -> ProblemOut:
    """标记一道错题为已复习：推进复习阶段并重算下次复习日期。"""
    today = date.today()
    conn = database.get_connection()
    try:
        row = _fetch_problem(conn, problem_id)
        if row is None:
            raise HTTPException(status_code=404, detail="错题不存在")

        new_stage = int(row["review_stage"]) + 1
        next_review = scheduler.next_review_date(new_stage, today)
        conn.execute(
            """
            UPDATE problems
            SET review_stage = ?, last_review_date = ?, next_review_date = ?
            WHERE id = ?
            """,
            (new_stage, today.isoformat(), next_review, problem_id),
        )
        conn.commit()
        row = _fetch_problem(conn, problem_id)
    finally:
        conn.close()
    return _row_to_problem(row)


# ---------------------------------------------------------------------------
# LaTeX 导出
# ---------------------------------------------------------------------------
@app.post("/api/export")
def export_latex(req: ExportRequest) -> Response:
    """按勾选的 problem_ids 生成 .tex 文件流供前端下载。"""
    if not req.problem_ids:
        raise HTTPException(status_code=400, detail="未选择任何错题")

    placeholders = ",".join("?" for _ in req.problem_ids)
    conn = database.get_connection()
    try:
        rows = conn.execute(
            f"SELECT * FROM problems WHERE id IN ({placeholders})",
            tuple(req.problem_ids),
        ).fetchall()
    finally:
        conn.close()
    if not rows:
        raise HTTPException(status_code=404, detail="所选错题不存在")

    # 按用户勾选顺序排列，拆分题目/答案。
    by_id = {row["id"]: row for row in rows}
    ordered = [by_id[i] for i in req.problem_ids if i in by_id]
    problems = []
    for r in ordered:
        question, split_answer = _split_q_a(r["latex_code"] or r["raw_text"] or "")
        problems.append(
            {
                "subject": r["subject"],
                "question_latex": question or None,
                "answer_latex": r["answer_latex"] or split_answer,
                "raw_text": r["raw_text"],
            }
        )

    tex = exporter.build_tex(
        problems,
        include_answers=req.include_answers,
        answers_last=req.answers_last,
        language=req.language,
    )
    return Response(
        content=tex,
        media_type="application/x-tex",
        headers={"Content-Disposition": "attachment; filename=mistakes.tex"},
    )


if __name__ == "__main__":
    # 直接运行本文件即可启动服务（Electron 主进程通过 spawn 调用）。
    uvicorn.run(app, host="127.0.0.1", port=8000)
